"""PPTX generation from PresentationSpec."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from .schema import (
    AgendaSlide,
    AnySlide,
    ClosingSlide,
    ContentSlide,
    PresentationSpec,
    TitleSlide,
    TwoColumnSlide,
)
from .themes import Theme, get_theme

# ---------------------------------------------------------------------------
# Layout constants (EMU — 914400 EMU = 1 inch)
# ---------------------------------------------------------------------------
SLIDE_W = 12192000   # 13.333 inches
SLIDE_H = 6858000    # 7.5 inches

MARGIN_L = int(0.55 * 914400)   # left/right margin
MARGIN_R = int(0.55 * 914400)
HEADER_H = int(1.30 * 914400)   # header bar height
ACCENT_H = int(0.07 * 914400)   # accent stripe height
CONTENT_TOP = HEADER_H + ACCENT_H + int(0.18 * 914400)
CONTENT_H = SLIDE_H - CONTENT_TOP - int(0.3 * 914400)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R


def _inch(n: float) -> int:
    return int(n * 914400)


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _blank_slide(prs: PptxPresentation):
    """Add a slide using the blank layout."""
    blank = prs.slide_layouts[6]
    return prs.slides.add_slide(blank)


def _set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_rect(slide, left: int, top: int, width: int, height: int,
              color: RGBColor, line: bool = False):
    shape = slide.shapes.add_shape(1, Emu(left), Emu(top), Emu(width), Emu(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color
    return shape


def _add_textbox(slide, left: int, top: int, width: int, height: int,
                 text: str, font_size: int, color: RGBColor,
                 bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT,
                 word_wrap: bool = True, font_name: str | None = None,
                 space_before: int = 0) -> None:
    txb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = txb.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = align
    if space_before:
        p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    if font_name:
        run.font.name = font_name


def _add_bullets(slide, left: int, top: int, width: int, height: int,
                 items: list[str], font_size: int, color: RGBColor,
                 marker: str = "●", font_name: str | None = None,
                 indent: int = 0) -> None:
    txb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = txb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(4 if i > 0 else 0)
        run = p.add_run()
        run.text = f"{marker}  {item}" if not indent else f"  {marker}  {item}"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        if font_name:
            run.font.name = font_name


# ---------------------------------------------------------------------------
# Shared layout elements
# ---------------------------------------------------------------------------

def _draw_header(slide, theme: Theme, title: str) -> None:
    """Draw the dark header bar + accent stripe + title text."""
    # Header background
    _add_rect(slide, 0, 0, SLIDE_W, HEADER_H, theme.header_bg)
    # Accent stripe below header
    _add_rect(slide, 0, HEADER_H, SLIDE_W, ACCENT_H, theme.accent)
    # Title text inside header
    _add_textbox(
        slide,
        left=MARGIN_L, top=int(0.32 * 914400),
        width=CONTENT_W, height=int(0.80 * 914400),
        text=title,
        font_size=theme.heading_size,
        color=theme.text_on_dark,
        bold=True,
        align=PP_ALIGN.LEFT,
    )


# ---------------------------------------------------------------------------
# Slide-type renderers
# ---------------------------------------------------------------------------

def _render_title(slide, spec: TitleSlide, theme: Theme) -> None:
    # Full-slide background
    _set_bg(slide, theme.header_bg)

    # Bottom accent bar
    accent_bar_h = _inch(0.45)
    _add_rect(slide, 0, SLIDE_H - accent_bar_h, SLIDE_W, accent_bar_h, theme.accent)

    # Left accent stripe (decorative)
    _add_rect(slide, 0, 0, _inch(0.12), SLIDE_H, theme.accent)

    # Title — vertically centred-ish
    title_top = _inch(2.1)
    _add_textbox(
        slide,
        left=_inch(0.9), top=title_top,
        width=_inch(11.5), height=_inch(1.6),
        text=spec.title,
        font_size=theme.title_size,
        color=theme.text_on_dark,
        bold=True,
        align=PP_ALIGN.LEFT,
    )

    # Subtitle
    if spec.subtitle:
        _add_textbox(
            slide,
            left=_inch(0.9), top=title_top + _inch(1.7),
            width=_inch(11.5), height=_inch(0.8),
            text=spec.subtitle,
            font_size=theme.body_size + 2,
            color=theme.text_muted,
            align=PP_ALIGN.LEFT,
        )


def _render_agenda(slide, spec: AgendaSlide, theme: Theme) -> None:
    _set_bg(slide, theme.slide_bg)
    _draw_header(slide, theme, spec.title)

    # Two-column layout for agenda items
    mid = len(spec.items) // 2 + len(spec.items) % 2
    left_items = spec.items[:mid]
    right_items = spec.items[mid:]

    col_w = (CONTENT_W - _inch(0.5)) // 2
    col_top = CONTENT_TOP
    col_h = CONTENT_H

    # Number + text
    for col_idx, col_items in enumerate([left_items, right_items]):
        col_left = MARGIN_L + col_idx * (col_w + _inch(0.5))
        for row_idx, item in enumerate(col_items):
            num = row_idx + 1 + col_idx * mid
            row_top = col_top + row_idx * _inch(0.85)

            # Number bubble (accent color)
            _add_textbox(
                slide,
                left=col_left, top=row_top,
                width=_inch(0.5), height=_inch(0.6),
                text=str(num),
                font_size=theme.bullet_size,
                color=theme.accent,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            # Item text
            _add_textbox(
                slide,
                left=col_left + _inch(0.6), top=row_top,
                width=col_w - _inch(0.65), height=_inch(0.6),
                text=item,
                font_size=theme.bullet_size,
                color=theme.text_on_light,
            )


def _render_content(slide, spec: ContentSlide, theme: Theme) -> None:
    _set_bg(slide, theme.slide_bg)
    _draw_header(slide, theme, spec.title)

    top = CONTENT_TOP

    if spec.body:
        _add_textbox(
            slide,
            left=MARGIN_L, top=top,
            width=CONTENT_W, height=_inch(1.0),
            text=spec.body,
            font_size=theme.body_size,
            color=theme.text_on_light,
        )
        top += _inch(1.1)

    if spec.bullets:
        _add_bullets(
            slide,
            left=MARGIN_L, top=top,
            width=CONTENT_W, height=SLIDE_H - top - _inch(0.3),
            items=spec.bullets,
            font_size=theme.bullet_size,
            color=theme.text_on_light,
            marker=theme.bullet_char,
        )


def _render_two_column(slide, spec: TwoColumnSlide, theme: Theme) -> None:
    _set_bg(slide, theme.slide_bg)
    _draw_header(slide, theme, spec.title)

    col_w = (CONTENT_W - _inch(0.4)) // 2
    top = CONTENT_TOP

    for col_idx, col in enumerate([spec.left, spec.right]):
        col_left = MARGIN_L + col_idx * (col_w + _inch(0.4))

        cur_top = top
        if col.title:
            _add_textbox(
                slide,
                left=col_left, top=cur_top,
                width=col_w, height=_inch(0.5),
                text=col.title,
                font_size=theme.body_size,
                color=theme.accent,
                bold=True,
            )
            cur_top += _inch(0.55)

        if col.body:
            _add_textbox(
                slide,
                left=col_left, top=cur_top,
                width=col_w, height=_inch(0.8),
                text=col.body,
                font_size=theme.bullet_size,
                color=theme.text_on_light,
            )
            cur_top += _inch(0.9)

        if col.bullets:
            _add_bullets(
                slide,
                left=col_left, top=cur_top,
                width=col_w, height=SLIDE_H - cur_top - _inch(0.3),
                items=col.bullets,
                font_size=theme.bullet_size - 1,
                color=theme.text_on_light,
                marker=theme.bullet_char,
            )

    # Vertical divider
    div_x = MARGIN_L + col_w + _inch(0.18)
    _add_rect(slide, div_x, CONTENT_TOP, _inch(0.02),
              CONTENT_H, theme.accent)


def _render_closing(slide, spec: ClosingSlide, theme: Theme) -> None:
    _set_bg(slide, theme.header_bg)

    # Top accent bar
    _add_rect(slide, 0, 0, SLIDE_W, _inch(0.45), theme.accent)
    # Bottom accent bar
    _add_rect(slide, 0, SLIDE_H - _inch(0.45), SLIDE_W, _inch(0.45), theme.accent)

    # Main title
    _add_textbox(
        slide,
        left=MARGIN_L, top=_inch(2.4),
        width=CONTENT_W, height=_inch(1.4),
        text=spec.title,
        font_size=theme.title_size - 4,
        color=theme.text_on_dark,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    if spec.message:
        _add_textbox(
            slide,
            left=MARGIN_L, top=_inch(4.0),
            width=CONTENT_W, height=_inch(0.8),
            text=spec.message,
            font_size=theme.body_size,
            color=theme.text_muted,
            align=PP_ALIGN.CENTER,
        )


# ---------------------------------------------------------------------------
# Dispatch table
# ---------------------------------------------------------------------------

_RENDERERS = {
    "title": _render_title,
    "agenda": _render_agenda,
    "content": _render_content,
    "two_column": _render_two_column,
    "closing": _render_closing,
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate(spec: PresentationSpec, output_path: str | Path) -> Path:
    """Generate a PPTX file from *spec* and save it to *output_path*."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    theme = get_theme(spec.theme)

    prs = PptxPresentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)

    for slide_spec in spec.slides:
        slide = _blank_slide(prs)
        renderer = _RENDERERS.get(slide_spec.type)
        if renderer:
            renderer(slide, slide_spec, theme)

    prs.save(str(output_path))
    return output_path
